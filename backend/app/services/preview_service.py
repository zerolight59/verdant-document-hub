"""Safe, bounded text previews; source files are never executed or sent to third parties."""

import csv
import io
from itertools import islice
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring
from fastapi import HTTPException

MIME_TYPES = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".txt": "text/plain",
    ".md": "text/plain",
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
PREVIEW_LIMIT = 20 * 1024 * 1024


def preview_content(path: Path) -> dict:
    suffix = path.suffix.lower()
    if not path.is_file():
        raise HTTPException(404, "Stored file is missing")
    if path.stat().st_size > PREVIEW_LIMIT:
        raise HTTPException(
            413, "Text previews are limited to 20 MB. Upload a PDF for larger documents."
        )
    if suffix in {".txt", ".md", ".csv"}:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        if suffix == ".csv":
            try:
                rows = list(islice(csv.reader(io.StringIO(text)), 2000))
            except csv.Error:
                raise HTTPException(
                    422, "This CSV is malformed or contains oversized cells."
                ) from None
            return {
                "sections": [{"title": "Data", "rows": [row[:50] for row in rows[:2000]]}],
                "note": "Preview shows up to 2,000 rows and 50 columns.",
            }
        return {
            "sections": [{"title": "", "text": text[:500000]}],
            "note": "Plain-text preview; up to 500,000 characters.",
        }
    if suffix not in {".docx", ".xlsx", ".pptx"}:
        raise HTTPException(
            415, "This format has no text preview. Use PDF, images, text or modern Office files."
        )
    try:
        with ZipFile(path) as archive:
            if (
                len(archive.infolist()) > 5000
                or sum(i.file_size for i in archive.infolist()) > PREVIEW_LIMIT
            ):
                raise HTTPException(
                    413, "This Office file is too complex for a safe preview. Use a PDF."
                )
            if suffix == ".xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                try:
                    sections = [
                        {
                            "title": sheet.title,
                            "rows": [
                                ["" if value is None else str(value) for value in row]
                                for row in sheet.iter_rows(
                                    max_row=min(sheet.max_row or 2000, 2000),
                                    max_col=min(sheet.max_column or 50, 50),
                                    values_only=True,
                                )
                            ],
                        }
                        for sheet in workbook.worksheets[:20]
                    ]
                finally:
                    workbook.close()
            else:
                names = (
                    ["word/document.xml"]
                    if suffix == ".docx"
                    else sorted(
                        [
                            name
                            for name in archive.namelist()
                            if name.startswith("ppt/slides/slide")
                            and name.endswith(".xml")
                            and "/_rels/" not in name
                        ],
                        key=lambda name: int(Path(name).stem.replace("slide", "")),
                    )
                )
                sections = []
                for index, name in enumerate(names[:200]):
                    root = fromstring(archive.read(name))
                    paragraphs = [
                        "".join(node.itertext()) for node in root.iter() if node.tag.endswith("}t")
                    ]
                    sections.append(
                        {
                            "title": f"Slide {index + 1}" if suffix == ".pptx" else "",
                            "text": "\n".join(paragraphs)[:500000],
                        }
                    )
        return {
            "sections": sections,
            "note": (
                "Readable content preview, not the original Office layout. Use PDF for exact "
                "formatting. Spreadsheets: up to 20 sheets, 2,000 rows and 50 columns."
            ),
        }
    except HTTPException:
        raise
    except (BadZipFile, KeyError, ValueError, OSError, ParseError, DefusedXmlException):
        raise HTTPException(
            422, "The file could not be previewed. Export it to PDF and upload a new version."
        ) from None
